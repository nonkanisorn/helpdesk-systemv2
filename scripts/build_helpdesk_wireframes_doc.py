from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


ROOT = Path("/Users/nonkns/Documents/GitHub/helpdesk-systemv2")
SRC_DOC = Path("/Users/nonkns/Documents/Codex/2026-05-14/files-mentioned-by-the-user-docx/โปรเจคระบบแจ้งซ่อม_แก้บทที่2.docx")
OUT_DOC = ROOT / "โปรเจคระบบแจ้งซ่อม_เพิ่ม_wireframe_sequence.docx"
ASSET_DIR = ROOT / "docs" / "generated-wireframes"
FONT_PATH = Path("/Library/Fonts/Arial Unicode.ttf")
if not FONT_PATH.exists():
    FONT_PATH = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")


def font(size, bold=False):
    return ImageFont.truetype(str(FONT_PATH), size=size)


def round_rect(draw, box, radius=18, outline="#9a9a9a", width=3, fill="#ffffff"):
    draw.rounded_rectangle(box, radius=radius, outline=outline, width=width, fill=fill)


def text(draw, xy, value, size=34, fill="#222222", anchor=None):
    draw.text(xy, value, font=font(size), fill=fill, anchor=anchor)


def line(draw, xy, fill="#555555", width=3):
    draw.line(xy, fill=fill, width=width)


def arrow(draw, start, end, label="", y_offset=-34):
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill="#333333", width=3)
    if x2 >= x1:
        pts = [(x2, y2), (x2 - 18, y2 - 9), (x2 - 18, y2 + 9)]
    else:
        pts = [(x2, y2), (x2 + 18, y2 - 9), (x2 + 18, y2 + 9)]
    draw.polygon(pts, fill="#333333")
    if label:
        text(draw, ((x1 + x2) / 2, y1 + y_offset), label, size=25, anchor="mm")


def nav(draw, w):
    text(draw, (42, 43), "HELPDESK APP", 35)
    for x, label in [(360, "หน้าหลัก"), (525, "แจ้งซ่อม"), (700, "ติดตามสถานะ/ประวัติของฉัน")]:
        text(draw, (x, 48), label, 28, fill="#555555")
    draw.ellipse((w - 95, 32, w - 55, 72), outline="#777777", width=4)
    line(draw, (0, 100, w, 100), "#bdbdbd", 3)


def save_login(path):
    w, h = 1800, 1050
    im = Image.new("RGB", (w, h), "#f4f6fa")
    d = ImageDraw.Draw(im)
    round_rect(d, (610, 250, 1190, 810), 26, "#c7c7c7", 3, "#ffffff")
    round_rect(d, (850, 305, 950, 405), 24, "#777777", 3, "#f2f2f2")
    text(d, (900, 470), "Helpdesk", 54, anchor="mm")
    text(d, (900, 525), "เข้าสู่ระบบเพื่อจัดการงานแจ้งซ่อม", 28, "#666666", anchor="mm")
    text(d, (675, 595), "Email", 24, "#555555")
    round_rect(d, (660, 620, 1140, 690), 8, "#777777", 3)
    text(d, (675, 735), "Password", 24, "#555555")
    round_rect(d, (660, 760, 1140, 830), 8, "#777777", 3)
    round_rect(d, (660, 875, 1140, 945), 8, "#555555", 3, "#e9e9e9")
    text(d, (900, 910), "Login", 34, anchor="mm")
    im.save(path)


def save_create_ticket(path):
    w, h = 1800, 1050
    im = Image.new("RGB", (w, h), "#f2f3f7")
    d = ImageDraw.Draw(im)
    nav(d, w)
    round_rect(d, (45, 150, 1755, 960), 20, "#c6c6c6", 3, "#ffffff")
    text(d, (95, 210), "แจ้งซ่อม", 44)
    text(d, (95, 265), "กรุณากรอกข้อมูลให้ครบถ้วน", 26, "#666666")
    fields = [
        (95, 335, 1705, 430, "หัวข้อการแจ้งซ่อม *"),
        (95, 515, 1705, 695, "รายละเอียดปัญหา *"),
        (95, 790, 1705, 880, "หมวดหมู่ปัญหา *"),
    ]
    for x1, y1, x2, y2, label in fields:
        round_rect(d, (x1, y1, x2, y2), 8, "#888888", 3)
        text(d, (x1 + 25, y1 + 38), label, 30, "#444444")
    round_rect(d, (1440, 910, 1580, 965), 8, "#777777", 3, "#ffffff")
    text(d, (1510, 938), "ล้างข้อมูล", 26, anchor="mm")
    round_rect(d, (1600, 910, 1725, 965), 8, "#777777", 3, "#eeeeee")
    text(d, (1662, 938), "เพิ่มการแจ้งซ่อม", 24, "#777777", anchor="mm")
    im.save(path)


def save_detail(path):
    w, h = 1800, 1120
    im = Image.new("RGB", (w, h), "#f2f3f7")
    d = ImageDraw.Draw(im)
    nav(d, w)
    round_rect(d, (20, 135, 1780, 245), 10, "#c6c6c6", 3, "#ffffff")
    text(d, (65, 175), "รายละเอียดการซ่อม", 38)
    round_rect(d, (35, 295, 1190, 690), 14, "#c6c6c6", 3, "#ffffff")
    text(d, (75, 345), "คอมเปิดไม่ติด", 36)
    round_rect(d, (80, 405, 610, 500), 8, "#9a9a9a", 2)
    text(d, (125, 440), "วันที่แจ้ง", 24, "#666666")
    round_rect(d, (635, 405, 1160, 500), 8, "#9a9a9a", 2)
    text(d, (680, 440), "ประเภทปัญหา", 24, "#666666")
    line(d, (75, 540, 1160, 540), "#cfcfcf", 2)
    text(d, (75, 585), "รายละเอียดปัญหา", 28)
    text(d, (75, 630), "ไม่ติดเลย", 26, "#555555")
    round_rect(d, (1215, 295, 1765, 470), 14, "#c6c6c6", 3, "#ffffff")
    text(d, (1255, 340), "ข้อมูลผู้แจ้ง", 28)
    round_rect(d, (1215, 505, 1765, 645), 14, "#c6c6c6", 3, "#ffffff")
    text(d, (1255, 550), "ข้อมูลผู้มอบหมาย", 28)
    round_rect(d, (1215, 680, 1765, 855), 14, "#c6c6c6", 3, "#ffffff")
    text(d, (1255, 725), "ข้อมูลอุปกรณ์", 28)
    round_rect(d, (35, 725, 1190, 1065), 14, "#c6c6c6", 3, "#ffffff")
    text(d, (75, 775), "ประวัติการดำเนินงาน", 34)
    for i, (label, y) in enumerate([("เปิดงาน", 840), ("มอบหมายแล้ว", 895), ("กำลังดำเนินการ", 950), ("เสร็จสิ้น", 1005)]):
        d.ellipse((75, y, 105, y + 30), outline="#777777", width=3)
        text(d, (125, y - 2), label, 24)
    im.save(path)


def save_history(path):
    w, h = 1800, 1050
    im = Image.new("RGB", (w, h), "#f2f3f7")
    d = ImageDraw.Draw(im)
    nav(d, w)
    round_rect(d, (35, 140, 1765, 420), 14, "#c6c6c6", 3, "#ffffff")
    text(d, (80, 205), "ประวัติการซ่อมของฉัน", 44)
    text(d, (80, 260), "รายการคำขอซ่อมทั้งหมดของคุณ", 28, "#666666")
    for i, label in enumerate(["ทั้งหมด", "รอดำเนินการ", "อยู่ระหว่างซ่อม", "รอการยืนยัน", "เสร็จสิ้น"]):
        x = 80 + i * 325
        round_rect(d, (x, 315, x + 250, 385), 10, "#999999", 3, "#ffffff")
        text(d, (x + 125, 350), label, 25, anchor="mm")
    round_rect(d, (35, 460, 1765, 985), 14, "#c6c6c6", 3, "#ffffff")
    text(d, (80, 535), "รายการแจ้งซ่อม", 42)
    round_rect(d, (80, 630, 875, 915), 12, "#999999", 3, "#ffffff")
    text(d, (120, 690), "คอมเปิดไม่ติด", 32)
    text(d, (120, 740), "ไม่ติดเลย", 26, "#666666")
    line(d, (120, 790, 835, 790), "#cccccc", 2)
    text(d, (120, 835), "แจ้งเมื่อ 15 พฤษภาคม 2569 เวลา 07:23", 25)
    round_rect(d, (120, 870, 835, 930), 8, "#777777", 3, "#ffffff")
    text(d, (477, 900), "ดูรายละเอียด", 28, anchor="mm")
    im.save(path)


def save_sequence(path, title, steps):
    w, h = 1800, 1050
    im = Image.new("RGB", (w, h), "#ffffff")
    d = ImageDraw.Draw(im)
    text(d, (70, 55), title, 42)
    participants = [("ผู้ใช้", 210), ("หน้าเว็บ", 600), ("ระบบ/API", 1010), ("ฐานข้อมูล", 1410)]
    for name, x in participants:
        round_rect(d, (x - 115, 120, x + 115, 175), 8, "#555555", 3, "#f5f5f5")
        text(d, (x, 148), name, 27, anchor="mm")
        line(d, (x, 175, x, 950), "#aaaaaa", 2)
    y = 245
    for src, dst, label in steps:
        x1 = participants[src][1]
        x2 = participants[dst][1]
        arrow(d, (x1, y), (x2, y), label)
        y += 86
    im.save(path)


def set_thai_run(run, size=16, bold=False):
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0


def add_text_paragraph_before(doc, ref_p, value="", align=WD_ALIGN_PARAGRAPH.LEFT, size=16, bold=False):
    p = ref_p.insert_paragraph_before()
    format_paragraph(p, align)
    if value:
        r = p.add_run(value)
        set_thai_run(r, size=size, bold=bold)
    return p


def add_picture_before(doc, ref_p, image_path, width_inches=6.02):
    p = doc.add_paragraph()
    format_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run()
    r.add_picture(str(image_path), width=Inches(width_inches))
    ref_p._p.addprevious(p._p)
    return p


def add_caption_before(doc, ref_p, caption):
    p = add_text_paragraph_before(doc, ref_p, caption, WD_ALIGN_PARAGRAPH.CENTER, 16, False)
    return p


def build_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "wf_login": ASSET_DIR / "wireframe_uc01_login.png",
        "wf_create": ASSET_DIR / "wireframe_uc02_create_ticket.png",
        "wf_detail": ASSET_DIR / "wireframe_uc06_ticket_detail.png",
        "wf_history": ASSET_DIR / "wireframe_uc07_ticket_history.png",
        "seq_login": ASSET_DIR / "sequence_uc01_login.png",
        "seq_create": ASSET_DIR / "sequence_uc02_create_ticket.png",
        "seq_detail": ASSET_DIR / "sequence_uc06_ticket_detail.png",
        "seq_history": ASSET_DIR / "sequence_uc07_ticket_history.png",
    }
    save_login(assets["wf_login"])
    save_create_ticket(assets["wf_create"])
    save_detail(assets["wf_detail"])
    save_history(assets["wf_history"])
    save_sequence(
        assets["seq_login"],
        "UC01: เข้าสู่ระบบ",
        [
            (0, 1, "กรอก Email และ Password"),
            (1, 2, "ส่งข้อมูลเข้าสู่ระบบ"),
            (2, 3, "ตรวจสอบบัญชีผู้ใช้"),
            (3, 2, "ส่งผลการตรวจสอบ"),
            (2, 1, "สร้างเซสชัน/แจ้งผล"),
            (1, 0, "เข้าสู่หน้าหลัก"),
        ],
    )
    save_sequence(
        assets["seq_create"],
        "UC02: เพิ่มการแจ้งซ่อม",
        [
            (0, 1, "เปิดหน้าแจ้งซ่อม"),
            (0, 1, "กรอกหัวข้อ รายละเอียด และหมวดหมู่"),
            (1, 2, "ส่งคำขอเพิ่มการแจ้งซ่อม"),
            (2, 3, "บันทึกข้อมูลคำร้อง"),
            (3, 2, "ส่งรหัสงานซ่อม"),
            (2, 1, "แสดงผลการบันทึก"),
            (1, 0, "แจ้งว่าสร้างรายการสำเร็จ"),
        ],
    )
    save_sequence(
        assets["seq_detail"],
        "UC06: ดูรายละเอียดการแจ้งซ่อม",
        [
            (0, 1, "กดปุ่มดูรายละเอียด"),
            (1, 2, "ขอข้อมูลรายการแจ้งซ่อม"),
            (2, 3, "ค้นหารายละเอียด/สถานะ/ผู้เกี่ยวข้อง"),
            (3, 2, "ส่งข้อมูลกลับ"),
            (2, 1, "จัดรูปแบบข้อมูลรายละเอียด"),
            (1, 0, "แสดงหน้ารายละเอียดการซ่อม"),
        ],
    )
    save_sequence(
        assets["seq_history"],
        "UC07: ดูรายการแจ้งซ่อมของตัวเอง",
        [
            (0, 1, "เปิดเมนูติดตามสถานะ/ประวัติของฉัน"),
            (1, 2, "ขอรายการแจ้งซ่อมของผู้ใช้"),
            (2, 3, "ค้นหารายการตามผู้แจ้ง"),
            (3, 2, "ส่งรายการและจำนวนสถานะ"),
            (2, 1, "แสดงสรุปและรายการแจ้งซ่อม"),
            (0, 1, "เลือกตัวกรองสถานะ"),
            (1, 0, "แสดงรายการตามตัวกรอง"),
        ],
    )
    return assets


def build_doc():
    assets = build_assets()
    doc = Document(str(SRC_DOC))
    ref = None
    for p in doc.paragraphs:
        if p.text.strip() == "บทที่ 4":
            ref = p
            break
    if ref is None:
        raise RuntimeError("Cannot find insertion point: บทที่ 4")

    add_text_paragraph_before(doc, ref, "3.8 การออกแบบหน้าจอระบบแจ้งซ่อมแบบ Wireframe", size=16, bold=True)
    add_text_paragraph_before(
        doc,
        ref,
        "การออกแบบหน้าจอระบบแจ้งซ่อมจัดทำเป็น wireframe เพื่อแสดงโครงสร้างส่วนติดต่อผู้ใช้จากหน้าจอหลักที่ผู้ใช้ใช้งาน ได้แก่ การเข้าสู่ระบบ การเพิ่มรายการแจ้งซ่อม การดูรายละเอียดการแจ้งซ่อม และการดูรายการแจ้งซ่อมของตนเอง",
        size=16,
    )
    for key, caption in [
        ("wf_login", "รูปที่ 3.84  Wireframe หน้าเข้าสู่ระบบ"),
        ("wf_create", "รูปที่ 3.85  Wireframe หน้าเพิ่มการแจ้งซ่อม"),
        ("wf_detail", "รูปที่ 3.86  Wireframe หน้าดูรายละเอียดการแจ้งซ่อม"),
        ("wf_history", "รูปที่ 3.87  Wireframe หน้าดูรายการแจ้งซ่อมของตัวเอง"),
    ]:
        add_picture_before(doc, ref, assets[key])
        add_caption_before(doc, ref, caption)

    add_text_paragraph_before(doc, ref, "3.9 แผนภาพแสดงลำดับการทำงานของระบบแจ้งซ่อม (Sequence Diagram)", size=16, bold=True)
    add_text_paragraph_before(
        doc,
        ref,
        "แผนภาพลำดับการทำงานต่อไปนี้อ้างอิงจากการกระทำของผู้ใช้ในหน้าจอที่จัดทำเป็น wireframe โดยแสดงการสื่อสารระหว่างผู้ใช้ หน้าเว็บ ระบบ/API และฐานข้อมูล",
        size=16,
    )
    for uc, key, caption in [
        ("UC01: เข้าสู่ระบบ", "seq_login", "รูปที่ 3.88  แผนภาพแสดงลำดับการทำงานการเข้าสู่ระบบ"),
        ("UC02: เพิ่มการแจ้งซ่อม", "seq_create", "รูปที่ 3.89  แผนภาพแสดงลำดับการทำงานการเพิ่มการแจ้งซ่อม"),
        ("UC06: ดูรายละเอียดการแจ้งซ่อม", "seq_detail", "รูปที่ 3.90  แผนภาพแสดงลำดับการทำงานการดูรายละเอียดการแจ้งซ่อม"),
        ("UC07: ดูรายการแจ้งซ่อมของตัวเอง", "seq_history", "รูปที่ 3.91  แผนภาพแสดงลำดับการทำงานการดูรายการแจ้งซ่อมของตัวเอง"),
    ]:
        add_text_paragraph_before(doc, ref, uc, size=16, bold=True)
        add_picture_before(doc, ref, assets[key])
        add_caption_before(doc, ref, caption)

    doc.save(str(OUT_DOC))
    print(OUT_DOC)


if __name__ == "__main__":
    build_doc()
